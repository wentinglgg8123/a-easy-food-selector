"""Validate the bundled workbook independently from its parsed restaurant objects."""

from openpyxl import load_workbook
import pytest

from food_picker import defaults
from food_picker.models import Filters, ValidationError
from food_picker.session import PickerSession


def test_bundled_catalog_preserves_all_original_rows_and_fields():
    workbook = load_workbook(defaults.DEFAULT_DATA_PATH, read_only=True, data_only=False)
    rows = [row for row in workbook.worksheets[0].iter_rows(max_col=11, values_only=True) if any(value is not None for value in row)]
    workbook.close()
    parsed = defaults.load_default_catalog().restaurants
    assert rows[0][:4] == ("Restaurant", "Location", "Gener", "Price")
    assert len(rows[1:]) == len(parsed) == 10
    for source, restaurant in zip(rows[1:], parsed, strict=True):
        assert restaurant.name == source[0]
        assert restaurant.location == source[1]
        assert restaurant.genre == source[2]
        assert restaurant.price is None
        assert restaurant.price_level == source[3]
        assert (
            restaurant.one_person, restaurant.dating, restaurant.gathering,
            restaurant.fast_serve, restaurant.slow_eat,
        ) == tuple(value == "O" for value in source[4:9])
        assert restaurant.latitude is None and restaurant.longitude is None
        assert restaurant.opening_hours == (source[9] or "")
        assert restaurant.special_opening_hours == (source[10] or "")
        assert restaurant.note == ""
    assert sum(restaurant.price_level == 1.5 for restaurant in parsed) == 5
    assert defaults.load_default_catalog().warnings == []


def test_current_source_rows_are_all_available_for_recommendations():
    result = defaults.load_default_catalog()
    closed = [restaurant for restaurant in result.restaurants if restaurant.is_closed]
    assert closed == []
    session = PickerSession()
    assert len(session.candidates(Filters())) == 10
    assert len(session.candidates(Filters(include_closed=True))) == 10


def test_each_default_load_returns_independent_mutable_collections():
    first = defaults.load_default_catalog()
    second = defaults.load_default_catalog()
    assert first.restaurants == second.restaurants
    assert first.restaurants is not second.restaurants
    assert first.warnings is not second.warnings
    first.restaurants.clear()
    first.warnings.append("只屬於第一個讀取結果")
    assert len(second.restaurants) == 10
    assert "只屬於第一個讀取結果" not in second.warnings


def test_missing_default_file_raises_explicit_error_without_demo_fallback(monkeypatch, tmp_path):
    monkeypatch.setattr(defaults, "DEFAULT_DATA_PATH", tmp_path / "missing.xlsx")
    with pytest.raises(ValidationError, match="內建餐廳名單.*存在"):
        defaults.load_default_catalog()
    with pytest.raises(ValidationError, match="內建餐廳名單"):
        PickerSession()


@pytest.mark.parametrize("content", [b"", b"not a spreadsheet"])
def test_corrupt_default_file_raises_explicit_error_without_demo_fallback(monkeypatch, tmp_path, content):
    path = tmp_path / "food.xlsx"
    path.write_bytes(content)
    monkeypatch.setattr(defaults, "DEFAULT_DATA_PATH", path)
    with pytest.raises(ValidationError, match="內建餐廳名單.*無法載入"):
        defaults.load_default_catalog()


def test_session_import_and_reset_do_not_change_bundled_workbook():
    before = defaults.DEFAULT_DATA_PATH.read_bytes()
    session = PickerSession()
    session.restaurants.clear()
    session.reset()
    assert len(session.restaurants) == 10
    assert defaults.DEFAULT_DATA_PATH.read_bytes() == before

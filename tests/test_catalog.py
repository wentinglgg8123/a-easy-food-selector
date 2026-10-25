"""Behavioral checks for import rules and recommendation boundaries."""

from dataclasses import FrozenInstanceError, replace
from datetime import date, datetime
from pathlib import Path
import random

import pytest

from food_picker.catalog import (
    choose_restaurant,
    demo_restaurants,
    filter_restaurants,
    genre_label,
    location_label,
    parse_rows,
    restaurant_categories,
)
from food_picker.models import Filters, ValidationError


def test_chinese_and_english_headers_with_optional_columns():
    result = parse_rows([
        [],
        ["Ｒｅｓｔａｕｒａｎｔ＿Ｎａｍｅ", "區域（校內/校外）", "Type", "均消（元）", "distance-km", "住址", "notes", "ignored"],
        [" 午餐店 ", "On Campus", "中式", "ＮＴ＄１，０００元", None, " 新竹市 ", "營業時間依現場為準", "自訂欄位"],
        ["早餐店", "off-campus", "", "$ 60", "２ km", "光復路", "", None],
    ])
    first, second = result.restaurants
    assert (first.name, first.area, first.category, first.price, first.address) == ("午餐店", "campus", "中式", 1_000, "新竹市")
    assert (second.area, second.category, second.distance) == ("nearby", "未分類", 2)
    assert result.warnings == []
    assert first.to_dict()["note"] == "營業時間依現場為準"
    with pytest.raises(FrozenInstanceError):
        first.name = "不能修改"


@pytest.mark.parametrize("value, expected", [(0, 0), (".5", .5), ("TWD 1,250.50", 1_250.5), ("80 NTD", 80), (None, None)])
def test_supported_price_formats(value, expected):
    result = parse_rows([["name", "area", "price"], ["店家", "campus", value]])
    assert result.restaurants[0].price == expected


@pytest.mark.parametrize("value", ["80-100", "1,00", "100人", "1e3", "-1", -1, float("inf"), float("nan"), True, date(2026, 10, 17), "=SUM(A1)", "\u3000＝1+2"])
def test_invalid_numeric_rows_are_skipped_without_losing_valid_rows(value):
    result = parse_rows([["name", "area", "price"], ["無效店家", "campus", value], ["有效店家", "campus", 80]])
    assert [restaurant.name for restaurant in result.restaurants] == ["有效店家"]
    assert len(result.warnings) == 1
    assert "第 2 列" in result.warnings[0]


def test_two_kilometre_boundary_and_missing_distance_warning():
    result = parse_rows([
        ["名稱", "地區", "距離"],
        ["範圍邊界", "校外", "2 公里"],
        ["範圍之外", "校外", 2.001],
        ["未填距離", "校外", ""],
        ["校內店家", "校內", 2.5],
    ])
    assert [restaurant.name for restaurant in result.restaurants] == ["範圍邊界", "未填距離", "校內店家"]
    assert any("第 3 列" in warning and "超過 2 公里" in warning for warning in result.warnings)
    assert any("第 4 列" in warning and "未填距離" in warning for warning in result.warnings)


def test_deduplication_uses_normalised_name_address_and_area():
    result = parse_rows([
        ["name", "area", "address"],
        ["Ｃａｆé", "campus", "Ａ路"],
        [" café ", "校內", "a路"],
        ["Café", "campus", "B路"],
        ["Café", "nearby", "Ａ路"],
    ])
    assert len(result.restaurants) == 3
    assert "第 3 列" in result.warnings[0] and "重複" in result.warnings[0]
    repeated = parse_rows([["name", "area", "address"], ["café", "campus", "a路"]])
    assert repeated.restaurants[0].id == result.restaurants[0].id


@pytest.mark.parametrize("rows, match", [
    ([], "沒有資料"),
    ([[None], [" "]], "沒有資料"),
    ([["name"], ["店家"]], "缺少必要欄位"),
    ([["name", "area", "店名"], ["店家", "campus", "店家"]], "出現多次"),
    ([["name", "area"], ["", "campus"]], "沒有可使用"),
])
def test_fatal_catalog_errors(rows, match):
    with pytest.raises(ValidationError, match=match):
        parse_rows(rows)


def test_limit_counts_nonempty_data_rows_and_aggregates_large_warning_groups():
    with pytest.raises(ValidationError, match="5,000"):
        parse_rows([["name", "area"]] + [["店家", "campus"]] * 5_001)
    result = parse_rows([["name", "area"]] + [["", "校內"]] * 14 + [[]] * 20 + [["店家", "校內"]])
    assert len(result.restaurants) == 1
    assert len(result.warnings) == 1
    assert "第 2、3、4、5、6、7、8、9、10、11、12、13 等 14 列" in result.warnings[0]


@pytest.mark.parametrize("column, size", [("name", 201), ("category", 81), ("address", 501), ("note", 2_001)])
def test_long_text_has_actionable_warning(column, size):
    headers = ["name", "area", "category", "address", "note"]
    invalid = ["店家", "campus", "中式", "地址", "備註"]
    invalid[headers.index(column)] = "文" * size
    result = parse_rows([headers, invalid, ["有效店家", "campus"]])
    assert len(result.restaurants) == 1
    assert "超過" in result.warnings[0] and "請縮短" in result.warnings[0]


@pytest.mark.parametrize("cell", ["=HYPERLINK(\"https://example.com\")", False, datetime(2026, 10, 17)])
def test_unsafe_cells_in_unknown_columns_are_rejected(cell):
    result = parse_rows([["name", "area", "extra"], ["無效店家", "campus", cell], ["有效店家", "campus", "自訂欄位"]])
    assert [restaurant.name for restaurant in result.restaurants] == ["有效店家"]


def test_filters_combine_and_unknown_price_is_not_within_budget():
    restaurants = demo_restaurants()
    filtered = filter_restaurants(restaurants, Filters(area="nearby", category="中式", max_price=100, query="巷口"))
    assert [restaurant.id for restaurant in filtered] == ["demo-nearby-dumplings"]
    assert filter_restaurants([replace(restaurants[0], price=None)], Filters(max_price=200)) == []
    assert filter_restaurants(restaurants, Filters(query="清華大學")) == restaurants
    normalised = replace(restaurants[0], name="ＣＡＦÉ", category="ＪＡＰＡＮＥＳＥ")
    assert filter_restaurants([normalised], Filters(query="café", category="japanese")) == [normalised]


def test_filter_defends_two_kilometre_rule_for_directly_created_records():
    restaurant = demo_restaurants()[4]
    assert filter_restaurants([replace(restaurant, distance=2)], Filters())
    for distance in (-.1, 2.01, float("nan"), float("inf")):
        assert filter_restaurants([replace(restaurant, distance=distance)], Filters()) == []


def test_selection_handles_empty_singleton_and_never_repeats_with_alternatives():
    restaurants = demo_restaurants()
    assert choose_restaurant([]) is None
    assert choose_restaurant(restaurants[:1], restaurants[0].id) == restaurants[0]
    rng = random.Random(2026)
    chosen_ids = {choose_restaurant(restaurants, restaurants[0].id, rng).id for _ in range(100)}
    assert restaurants[0].id not in chosen_ids
    assert chosen_ids == {restaurant.id for restaurant in restaurants[1:]}
    assert choose_restaurant(restaurants, rng=random.Random(1)) == choose_restaurant(restaurants, rng=random.Random(1))


def test_demo_catalog_is_fresh_and_explicitly_fictional():
    restaurants = demo_restaurants()
    assert len(restaurants) == 8
    assert len({restaurant.id for restaurant in restaurants}) == 8
    assert all("示範" in restaurant.name and "非真實店家名單" in restaurant.note for restaurant in restaurants)
    restaurants.clear()
    assert len(demo_restaurants()) == 8


RICH_HEADERS = [
    "Restaurant", "Location", "Genre", "Price", "OnePerson", "Dating",
    "Gathering", "FastServe", "SlowEat", "Coordinates", "Comments",
]


def test_rich_schema_retains_all_fields_and_keeps_tiers_separate_from_currency():
    result = parse_rows([
        RICH_HEADERS,
        ["牛排館", "Gongguan", "Steak", 4.0, "O", True, "○", "✓", 1, "25.018, 121.531", "學生優惠"],
    ])
    restaurant = result.restaurants[0]
    assert restaurant.area == "other"
    assert (restaurant.location, restaurant.genre, restaurant.category) == ("Gongguan", "Steak", "牛排")
    assert restaurant.price_level == 4 and restaurant.price is None
    assert all((restaurant.one_person, restaurant.dating, restaurant.gathering, restaurant.fast_serve, restaurant.slow_eat))
    assert (restaurant.latitude, restaurant.longitude) == (25.018, 121.531)
    assert restaurant.note == "學生優惠" and not restaurant.is_closed
    assert result.warnings == []


@pytest.mark.parametrize("value, expected", [(1, 1), (1.5, 1.5), (2.5, 2.5), (5.0, 5), ("３", 3), ("2.0", 2), ("１．５", 1.5), (None, None), ("", None)])
def test_price_levels_are_optional_half_steps_from_one_to_five(value, expected):
    restaurant = parse_rows([RICH_HEADERS, ["餐廳", "School", "Bento", value]]).restaurants[0]
    assert restaurant.price_level == expected
    assert restaurant.price is None


@pytest.mark.parametrize("value", [0, 6, 1.25, True, "NT$ 2", "1e0", "-1", "2-3"])
def test_invalid_price_level_skips_only_the_affected_restaurant(value):
    result = parse_rows([RICH_HEADERS, ["無效", "Gongguan", "Ramen", value], ["有效", "Gongguan", "Ramen", 2]])
    assert [restaurant.name for restaurant in result.restaurants] == ["有效"]
    assert "第 2 列" in result.warnings[0]


@pytest.mark.parametrize("value", ["bad coordinates", "91,121", "25,181", "nan,121", "25,121,22", 123])
def test_invalid_coordinates_warn_but_preserve_the_restaurant(value):
    row = ["餐廳", "Gongguan", "Ramen", 2, None, None, None, None, None, value]
    result = parse_rows([RICH_HEADERS, row])
    assert len(result.restaurants) == 1
    assert result.restaurants[0].latitude is None
    assert result.restaurants[0].longitude is None
    assert len(result.warnings) == 1 and "已保留餐廳" in result.warnings[0]


def test_blank_coordinates_are_allowed_without_a_warning():
    result = parse_rows([RICH_HEADERS, ["餐廳", "School", "Bento", None]])
    assert result.restaurants[0].latitude is None
    assert result.warnings == []


def test_rich_schema_validates_flags_and_does_not_allow_booleans_in_other_columns():
    result = parse_rows([
        RICH_HEADERS,
        ["一人食", "Gongguan", "Ramen", 1, "yes", "false", 0, None, False],
        ["無效標記", "Gongguan", "Ramen", 1, "maybe"],
        ["布林名稱", "Gongguan", True, 1],
    ])
    restaurant, = result.restaurants
    assert restaurant.one_person and not restaurant.dating and not restaurant.gathering
    assert not restaurant.fast_serve and not restaurant.slow_eat
    assert len(result.warnings) == 2
    assert any("OnePerson" in warning for warning in result.warnings)


def test_rich_schema_rejects_formulas_in_coordinates_and_comments():
    result = parse_rows([
        RICH_HEADERS,
        ["座標公式", "Gongguan", "Ramen", 2, None, None, None, None, None, "=A1"],
        ["備註公式", "Gongguan", "Ramen", 2, None, None, None, None, None, None, "=A1"],
        ["有效餐廳", "Gongguan", "Ramen", 2],
    ])
    assert [restaurant.name for restaurant in result.restaurants] == ["有效餐廳"]
    assert len(result.warnings) == 1 and "第 2、3 列" in result.warnings[0]


def test_rich_schema_requires_location_and_genre_headers():
    with pytest.raises(ValidationError, match="Genre"):
        parse_rows([["Restaurant", "Location"], ["餐廳", "Gongguan"]])
    with pytest.raises(ValidationError, match="Location"):
        parse_rows([["Restaurant", "Genre"], ["餐廳", "Ramen"]])


def test_unknown_codes_fall_back_without_inventing_a_location_or_cuisine():
    restaurant = parse_rows([RICH_HEADERS, ["餐廳", "NewTown", "NewCuisine", 2]]).restaurants[0]
    assert restaurant.location == "NewTown" and restaurant.category == "NewCuisine"
    assert location_label(restaurant.location) == "NewTown"
    assert genre_label(restaurant.genre) == "NewCuisine"
    assert genre_label("") == "未分類"


def test_rich_deduplication_keeps_distinct_branches():
    rows = [
        ["同名店", "Gongguan", "Ramen", 1, None, None, None, None, None, "25.01,121.53"],
        ["同名店", "Gongguan", "Ramen", 1, None, None, None, None, None, "25.02,121.53"],
        ["同名店", "Shida", "Ramen", 1, None, None, None, None, None, "25.01,121.53"],
    ]
    result = parse_rows([RICH_HEADERS, *rows, rows[0]])
    assert len(result.restaurants) == 3
    assert len({restaurant.id for restaurant in result.restaurants}) == 3
    assert len(result.warnings) == 1 and "重複" in result.warnings[0]


def test_rich_schema_can_round_trip_optional_legacy_currency_and_distance_fields():
    result = parse_rows([
        [*RICH_HEADERS, "均消", "距離（公里）", "地址", "區域"],
        ["校內店家", "", "中式", None, None, None, None, None, None, None, "原始備註", 100, None, "清大校內", "campus"],
    ])
    restaurant = result.restaurants[0]
    assert restaurant.area == "campus" and restaurant.price == 100
    assert restaurant.price_level is None and restaurant.address == "清大校內"
    assert restaurant.category == "中式" and restaurant.note == "原始備註"


def test_new_filters_combine_and_closed_restaurants_are_opt_in():
    restaurants = parse_rows([
        RICH_HEADERS,
        ["全部符合", "Gongguan", "Ramen", 2, "O", "O", "O", "O", "O"],
        ["少一條件", "Gongguan", "Ramen", 2, "O", "O", "O", "O", None],
        ["已關門", "Gongguan", "Ramen", 1, "O", "O", "O", "O", "O", None, "關門了"],
        ["未知價位", "Gongguan", "Ramen", None, "O", "O", "O", "O", "O"],
        ["不同地點", "Shida", "Ramen", 1, "O", "O", "O", "O", "O"],
    ]).restaurants
    filters = Filters(location="Gongguan", category="Ramen", max_price_level=2,
                      one_person=True, dating=True, gathering=True, fast_serve=True, slow_eat=True)
    assert [restaurant.name for restaurant in filter_restaurants(restaurants, filters)] == ["全部符合"]
    included = filter_restaurants(restaurants, replace(filters, include_closed=True))
    assert [restaurant.name for restaurant in included] == ["全部符合", "已關門"]
    assert len(filter_restaurants(restaurants, Filters(query="公館"))) == 3
    assert len(filter_restaurants(restaurants, Filters(query="gongguan"))) == 3
    assert len(filter_restaurants(restaurants, Filters(query="拉麵"))) == 4
    assert filter_restaurants(restaurants, Filters(max_price=100)) == []


@pytest.fixture
def multiselect_catalog():
    return parse_rows([
        RICH_HEADERS,
        ["公館拉麵", "Gongguan", "Ramen", 1, "O", None, "O"],
        ["師大甜點", "Shida", "Dessert", 3, "O", None, "O"],
        ["新生拉麵", "XinshengSouth", "Ramen", 2, "O", None, "O"],
        ["公館牛排", "Gongguan", "Steak", 3, "O", None, None],
        ["公館未知價位", "Gongguan", "Ramen", None, "O", None, "O"],
        ["師大歇業", "Shida", "Dessert", 1, "O", None, "O", None, None, None, "關門了"],
    ]).restaurants


def test_location_multiselect_matches_a_union_without_duplicating_restaurants(multiselect_catalog):
    filters = Filters(locations=("Ｇｏｎｇｇｕａｎ", "shida", "Gongguan"))
    assert [restaurant.name for restaurant in filter_restaurants(multiselect_catalog, filters)] == [
        "公館拉麵", "師大甜點", "公館牛排", "公館未知價位",
    ]


def test_category_multiselect_matches_raw_codes_and_localized_labels(multiselect_catalog):
    filters = Filters(categories=("ＲＡＭＥＮ", "甜點"))
    assert [restaurant.name for restaurant in filter_restaurants(multiselect_catalog, filters)] == [
        "公館拉麵", "師大甜點", "新生拉麵", "公館未知價位",
    ]


def test_price_multiselect_matches_integer_tiers_and_excludes_unknowns(multiselect_catalog):
    filtered = filter_restaurants(multiselect_catalog, Filters(price_levels=(1, 3)))
    assert [restaurant.name for restaurant in filtered] == ["公館拉麵", "師大甜點", "公館牛排"]
    assert "新生拉麵" not in {restaurant.name for restaurant in filtered}
    unrestricted = filter_restaurants(multiselect_catalog, Filters(price_levels=()))
    assert "公館未知價位" in {restaurant.name for restaurant in unrestricted}


@pytest.fixture
def half_price_catalog():
    return parse_rows([
        RICH_HEADERS,
        *[[f"價格{level}店", "小吃部", "便當", level]
          for level in (1, 1.5, 2, 2.5, 3, 3.5, 4, 4.5, 5)],
        ["未知價格店", "小吃部", "便當", None],
    ]).restaurants


@pytest.mark.parametrize("button, expected", [
    (1, [1, 1.5]),
    (2, [1.5, 2, 2.5]),
    (3, [2.5, 3, 3.5]),
    (4, [3.5, 4, 4.5]),
    (5, [4.5, 5]),
])
def test_integer_price_buttons_include_only_their_own_tier_and_adjacent_half_tiers(half_price_catalog, button, expected):
    filtered = filter_restaurants(half_price_catalog, Filters(price_levels=(button,)))
    assert [restaurant.price_level for restaurant in filtered] == expected
    assert all(restaurant.price is None for restaurant in filtered)
    assert [restaurant.price_level for restaurant in half_price_catalog] == [
        1, 1.5, 2, 2.5, 3, 3.5, 4, 4.5, 5, None,
    ]


def test_price_button_union_does_not_duplicate_half_tiers_or_include_unknowns(half_price_catalog):
    filtered = filter_restaurants(half_price_catalog, Filters(price_levels=(1, 2, 2)))
    assert [restaurant.price_level for restaurant in filtered] == [1, 1.5, 2, 2.5]
    assert len({restaurant.id for restaurant in filtered}) == len(filtered)
    all_known = filter_restaurants(half_price_catalog, Filters(price_levels=(1, 2, 3, 4, 5)))
    assert all_known == half_price_catalog[:-1]
    assert filter_restaurants(half_price_catalog, Filters()) == half_price_catalog
    # Existing programmatic callers can still request an exact fractional source value.
    assert [restaurant.price_level for restaurant in filter_restaurants(
        half_price_catalog, Filters(price_levels=(1.5,))
    )] == [1.5]


def test_multiselect_groups_and_flags_are_applied_together(multiselect_catalog):
    filters = Filters(
        locations=("Gongguan", "Shida"), categories=("Ramen", "Dessert", "Steak"),
        price_levels=(1, 3), one_person=True, gathering=True,
    )
    assert [restaurant.name for restaurant in filter_restaurants(multiselect_catalog, filters)] == [
        "公館拉麵", "師大甜點",
    ]
    assert [restaurant.name for restaurant in filter_restaurants(
        multiselect_catalog, replace(filters, include_closed=True)
    )] == ["公館拉麵", "師大甜點", "師大歇業"]


def test_legacy_singular_filters_and_maxima_remain_additional_constraints(multiselect_catalog):
    filters = Filters(
        location="Gongguan", locations=("Gongguan", "Shida"), category="Ramen",
        categories=("Ramen", "Dessert"), max_price_level=2, price_levels=(1, 3),
    )
    assert [restaurant.name for restaurant in filter_restaurants(multiselect_catalog, filters)] == ["公館拉麵"]
    assert filter_restaurants(multiselect_catalog, replace(filters, location="XinshengSouth")) == []
    assert filter_restaurants(multiselect_catalog, replace(filters, category="Steak")) == []
    assert filter_restaurants(multiselect_catalog, replace(filters, max_price=100)) == []


def test_empty_multiselect_groups_preserve_legacy_defaults_and_currency_budgets(multiselect_catalog):
    assert filter_restaurants(multiselect_catalog, Filters()) == filter_restaurants(
        multiselect_catalog, Filters(locations=(), categories=(), price_levels=())
    )
    demo = demo_restaurants()
    legacy_filters = Filters(area="campus", category="中式", max_price=100)
    assert filter_restaurants(demo, legacy_filters) == [demo[0]]
    assert filter_restaurants(demo, replace(legacy_filters, categories=("中式", "日式"))) == [demo[0]]
    assert filter_restaurants(demo, replace(legacy_filters, price_levels=(1, 2))) == []


def test_bundled_workbook_retains_all_rows_codes_and_fields():
    from openpyxl import load_workbook

    path = Path(__file__).resolve().parents[1] / "data" / "food.xlsx"
    if not path.is_file():
        pytest.skip("The user-supplied workbook has not been copied into data/food.xlsx yet")
    with path.open("rb") as source:
        workbook = load_workbook(source, read_only=True, data_only=False)
        try:
            result = parse_rows(list(workbook.active.iter_rows(values_only=True)))
        finally:
            workbook.close()
    restaurants = result.restaurants
    assert len(restaurants) == 10
    assert {restaurant.location for restaurant in restaurants} == {"小吃部", "清夜"}
    assert len({category for restaurant in restaurants for category in restaurant_categories(restaurant)}) == 18
    assert all(restaurant.area == "other" and restaurant.price is None for restaurant in restaurants)
    assert all(restaurant.latitude is None and restaurant.longitude is None for restaurant in restaurants)
    assert sum(restaurant.price_level == 1.5 for restaurant in restaurants) == 5
    assert sum(restaurant.one_person for restaurant in restaurants) == 9
    assert sum(restaurant.dating for restaurant in restaurants) == 0
    assert sum(restaurant.gathering for restaurant in restaurants) == 2
    assert sum(restaurant.fast_serve for restaurant in restaurants) == 9
    assert sum(restaurant.slow_eat for restaurant in restaurants) == 1
    assert sum(bool(restaurant.opening_hours) for restaurant in restaurants) == 9
    assert sum(bool(restaurant.special_opening_hours) for restaurant in restaurants) == 2
    closed = [restaurant for restaurant in restaurants if restaurant.is_closed]
    assert closed == []
    assert len(filter_restaurants(restaurants, Filters())) == 10
    assert result.warnings == []


def test_multi_genre_rows_accept_the_uploaded_header_and_filter_by_each_individual_cuisine():
    result = parse_rows([
        ["Restaurant", "Location", "Gener", "Price", "Available time", "special available time"],
        ["早餐店", "小吃部", "早餐, 早午餐，中式料理", 1.5, "7:30~19:30", "假日8:30~14:30"],
        ["速食店", "小吃部", "美式料理, 炸物, 速食", 2, "7:00~24:00", None],
    ])
    breakfast, fastfood = result.restaurants
    assert breakfast.genre == "早餐, 早午餐，中式料理"
    assert restaurant_categories(breakfast) == ("早餐", "早午餐", "中式料理")
    assert breakfast.price_level == 1.5 and breakfast.price is None
    assert breakfast.opening_hours == "7:30~19:30"
    assert breakfast.special_opening_hours == "假日8:30~14:30"
    assert filter_restaurants(result.restaurants, Filters(categories=("早餐",))) == [breakfast]
    assert filter_restaurants(result.restaurants, Filters(category="中式料理")) == [breakfast]
    assert filter_restaurants(result.restaurants, Filters(categories=("中式料理", "速食"))) == [breakfast, fastfood]
    assert filter_restaurants(result.restaurants, Filters(price_levels=(1.5,))) == [breakfast]
    assert filter_restaurants(result.restaurants, Filters(categories=("速食",), price_levels=(1.5,))) == []
    assert result.warnings == []

"""Real Chromium checks for shared restaurant data on desktop and mobile screens."""

from io import BytesIO
import os
from pathlib import Path
import re
import subprocess
import sys
import time
from urllib.error import URLError
from urllib.parse import parse_qs, urlparse
from urllib.request import urlopen

from openpyxl import Workbook, load_workbook
from playwright.sync_api import expect, sync_playwright
import pytest

from food_picker.excel import EXCEL_HEADERS

pytestmark = pytest.mark.browser
ROOT = Path(__file__).resolve().parents[1]
APP_URL = "http://127.0.0.1:8501"
MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
RESULT = 'section[aria-label="餐廳推薦結果"]'


def _server_ready() -> bool:
    try:
        with urlopen(f"{APP_URL}/_stcore/health", timeout=1) as response:
            return response.status == 200 and response.read() == b"ok"
    except (OSError, URLError):
        return False


@pytest.fixture(scope="module")
def app_url(tmp_path_factory):
    """Reuse a running local app, or start and clean up this fixture's own server."""
    if _server_ready():
        yield APP_URL
        return
    log_path = tmp_path_factory.mktemp("streamlit") / "server.log"
    with log_path.open("w+") as log:
        process = subprocess.Popen(
            [
                sys.executable, "-m", "streamlit", "run", "app.py",
                "--server.port", "8501", "--server.headless", "true",
                "--browser.gatherUsageStats", "false",
            ],
            cwd=ROOT,
            stdout=log,
            stderr=subprocess.STDOUT,
        )
        try:
            deadline = time.monotonic() + 30
            while not _server_ready():
                if process.poll() is not None or time.monotonic() >= deadline:
                    log.seek(0)
                    pytest.fail(f"Streamlit did not become ready:\n{log.read()}")
                time.sleep(0.2)
            yield APP_URL
        finally:
            process.terminate()
            try:
                process.wait(timeout=10)
            except subprocess.TimeoutExpired:
                process.kill()
                process.wait(timeout=5)


@pytest.fixture(scope="module")
def browser():
    with sync_playwright() as playwright:
        options = {"headless": True, "args": ["--no-sandbox"]}
        if executable := os.environ.get("CHROMIUM_PATH"):
            options["executable_path"] = executable
        instance = playwright.chromium.launch(**options)
        yield instance
        instance.close()


@pytest.fixture(params=[{"width": 1440, "height": 1000}, {"width": 390, "height": 844}], ids=["desktop", "mobile"])
def page(browser, app_url, request):
    context = browser.new_context(viewport=request.param, accept_downloads=True)
    page = context.new_page()
    page.set_default_timeout(15_000)
    errors = []
    page.on("pageerror", lambda error: errors.append(str(error)))
    page.goto(app_url)
    expect(page.get_by_role("button", name="幫我決定這一餐", exact=True)).to_be_visible(timeout=15_000)
    _wait_idle(page)
    try:
        yield page
        assert errors == [], f"Uncaught browser errors: {errors}"
        widths = page.evaluate("""() => ({
            viewport: window.innerWidth,
            document: document.documentElement.scrollWidth,
            main: document.querySelector('[data-testid="stMain"]').scrollWidth,
            mainClient: document.querySelector('[data-testid="stMain"]').clientWidth
        })""")
        assert widths["document"] <= widths["viewport"] + 1, widths
        assert widths["main"] <= widths["mainClient"] + 1, widths
    finally:
        context.close()


def _workbook(rows, *, extra_sheet=False) -> bytes:
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    if extra_sheet:
        second = workbook.create_sheet("第二張不應載入")
        second.append(EXCEL_HEADERS)
        second.append(["第二張不應載入", "Gongguan", "Bento", 1])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _catalog_bytes() -> bytes:
    return _workbook([
        EXCEL_HEADERS,
        ["測試旗標完整店", "Gongguan", "Bento", 2, "O", "O", "O", "O", "O", "25.01826359388706, 121.5310917601049", "測試學生優惠"],
        ["測試缺少節奏店", "Gongguan", "Bento", 2, "O", "O", "O", "", "O", "25.04, 121.55", "第二間測試店家"],
        ["測試別區店", "Shida", "Steak", 4, "", "O", "O", "", "O", "25.05, 121.56", ""],
    ], extra_sheet=True)


def _upload(page, content: bytes, filename="測試餐廳.xlsx"):
    page.get_by_role("button", name="匯入 Excel 名單", exact=True).click()
    dialog = page.get_by_role("dialog")
    expect(dialog).to_be_visible()
    dialog.locator('input[type="file"]').set_input_files(
        {"name": filename, "mimeType": MIME_XLSX, "buffer": content}
    )
    confirm = dialog.get_by_role("button", name="確認匯入", exact=True)
    expect(confirm).to_be_enabled()
    confirm.click()
    return dialog


def _import_catalog(page):
    dialog = _upload(page, _catalog_bytes())
    expect(dialog).not_to_be_visible()
    expect(page.get_by_text("已匯入 3 間餐廳，開始你的美食探索。", exact=True)).to_be_visible()
    _expect_count(page, 3)


def _select_pills(page, label, values):
    for value in values:
        _wait_idle(page)
        button = page.get_by_role("toolbar", name=label, exact=True).get_by_role(
            "button", name=value, exact=True
        )
        expect(button).to_have_attribute("aria-pressed", "false")
        button.click()
        expect(button).to_have_attribute("aria-pressed", "true")
        _wait_idle(page)


def _toggle_pill_off(page, label, value):
    _wait_idle(page)
    button = page.get_by_role("toolbar", name=label, exact=True).get_by_role(
        "button", name=value, exact=True
    )
    expect(button).to_have_attribute("aria-pressed", "true")
    button.click()
    expect(button).to_have_attribute("aria-pressed", "false")
    _wait_idle(page)


def _wait_idle(page):
    expect(page.get_by_test_id("stApp")).to_have_attribute("data-test-script-state", "notRunning")


def _query(page, value):
    query = page.get_by_role("textbox", name="搜尋店家", exact=True)
    query.fill(value)
    query.press("Enter")
    return query


def _expect_count(page, count, available=None):
    available = count if available is None else available
    _wait_idle(page)
    counter = page.get_by_text(f"{count} 間符合條件 · {available} 間可抽選", exact=True)
    expect(counter).to_have_count(1)
    expect(counter).to_be_visible()


def _download_rows(page, button_name, destination):
    with page.expect_download() as event:
        page.get_by_role("button", name=button_name, exact=True).click()
    event.value.save_as(destination)
    workbook = load_workbook(destination, read_only=True)
    rows = list(workbook.active.values)
    workbook.close()
    return rows


def test_default_catalog_is_shared_and_random_recommendations_do_not_repeat(page):
    expect(page.get_by_text("FOODFLIX", exact=True)).to_have_count(0)
    expect(page.get_by_role("combobox")).to_have_count(0)
    expect(page.get_by_role("toolbar")).to_have_count(5)
    expect(page.get_by_role("toolbar", name="食物種類", exact=True).get_by_role("button")).to_have_count(18)
    search_bounds = page.get_by_role("textbox", name="搜尋店家", exact=True).bounding_box()
    filter_bounds = page.locator(".st-key-filters_panel").bounding_box()
    assert search_bounds is not None and filter_bounds is not None
    assert 0 <= search_bounds["y"] < filter_bounds["y"] < page.viewport_size["height"]
    for label in ("地點", "食物種類"):
        bounds = page.get_by_role("toolbar", name=label, exact=True).get_by_role("button").first.bounding_box()
        assert bounds is not None and 0 <= bounds["y"] < page.viewport_size["height"], bounds
    _expect_count(page, 10)
    expect(page.get_by_role("button", name="下載完整名單", exact=True)).to_be_visible()
    page.get_by_role("button", name="幫我決定這一餐", exact=True).click()
    result = page.locator(f"{RESULT} h2")
    expect(result).not_to_have_text(re.compile(r"^示範・"))
    previous = result.inner_text()
    assert previous in {"包子商行", "麥當勞", "杜蘭小麥 - 義大利麵.披薩", "比司多早午餐", "7-ELEVEN", "漢城異國美食", "蘇記牛肉麵", "一番快餐", "好富熟成黑咖哩.台式蓋飯", "珈名鮮茶"}
    for _ in range(3):
        page.get_by_role("button", name="再推薦一家", exact=True).click()
        expect(result).not_to_have_text(previous)
        previous = result.inner_text()
        assert previous in {"包子商行", "麥當勞", "杜蘭小麥 - 義大利麵.披薩", "比司多早午餐", "7-ELEVEN", "漢城異國美食", "蘇記牛肉麵", "一番快餐", "好富熟成黑咖哩.台式蓋飯", "珈名鮮茶"}
    page.locator("summary").filter(has_text="最近推薦").click()
    expect(page.get_by_role("heading", name="最近推薦", exact=True)).to_be_visible()


def test_actual_location_genre_price_and_occasion_filters_then_reset(page):
    _select_pills(page, "地點", ["小吃部"])
    _expect_count(page, 9)
    _select_pills(page, "食物種類", ["速食"])
    _expect_count(page, 1)
    _select_pills(page, "價格等級", ["🪙🪙"])
    _select_pills(page, "用餐情境", ["一個人吃", "聚餐"])
    _expect_count(page, 1)
    page.get_by_role("button", name="幫我決定這一餐", exact=True).click()
    expect(page.locator(f"{RESULT} h2")).to_have_text("麥當勞")
    expect(page.locator(RESULT)).to_contain_text("🪙🪙 · NT$100–200")
    expect(page.locator(RESULT)).to_contain_text("7:00~24:00")
    query = _query(page, "不存在的餐廳")
    _expect_count(page, 0)
    expect(page.get_by_role("button", name="幫我決定這一餐", exact=True)).to_be_disabled()
    page.get_by_role("button", name="全部取消", exact=True).click()
    _expect_count(page, 10)
    expect(query).to_have_value("")
    expect(page.get_by_role("toolbar", name="食物種類", exact=True).get_by_role("button", pressed=True)).to_have_count(0)
    expect(page.get_by_role("toolbar", name="價格等級", exact=True).get_by_role("button", pressed=True)).to_have_count(0)
    for label in ("地點", "用餐情境", "用餐節奏"):
        expect(page.get_by_role("toolbar", name=label, exact=True).get_by_role("button", pressed=True)).to_have_count(0)


def test_rich_import_requires_all_features_uses_named_maps_and_restores_defaults(page):
    _import_catalog(page)
    _select_pills(page, "地點", ["公館"])
    _select_pills(page, "食物種類", ["便當"])
    _select_pills(page, "價格等級", ["🪙🪙"])
    _select_pills(page, "用餐情境", ["一個人吃", "約會", "聚餐"])
    _select_pills(page, "用餐節奏", ["快速上菜", "慢慢享用"])
    _expect_count(page, 1)
    page.get_by_role("button", name="幫我決定這一餐", exact=True).click()
    expect(page.locator(f"{RESULT} h2")).to_have_text("測試旗標完整店")
    expect(page.locator(RESULT)).to_contain_text("測試學生優惠")
    link = page.get_by_role("link", name="在 Google 地圖查看", exact=True)
    expect(link).to_be_visible()
    query = parse_qs(urlparse(link.get_attribute("href")).query)["query"][0]
    assert "測試旗標完整店" in query and "公館" in query
    page.get_by_role("button", name="再推薦一家", exact=True).click()
    expect(page.locator(f"{RESULT} h2")).to_have_text("測試旗標完整店")
    page.get_by_role("button", name="恢復內建名單", exact=True).click()
    _expect_count(page, 10)
    expect(page.get_by_role("button", name="恢復內建名單", exact=True)).not_to_be_visible()
    expect(page.locator(RESULT)).not_to_be_visible()
    expect(page.get_by_role("toolbar", name="食物種類", exact=True).get_by_role("button", pressed=True)).to_have_count(0)


def test_invalid_import_preserves_the_current_custom_catalog(page):
    _import_catalog(page)
    dialog = _upload(page, _workbook([["錯誤欄位"], ["沒有必要表頭"]]), "錯誤欄位.xlsx")
    expect(dialog.get_by_text(re.compile("缺少必要欄位"))).to_be_visible()
    expect(dialog).to_be_visible()
    page.keyboard.press("Escape")
    expect(dialog).not_to_be_visible()
    _expect_count(page, 3)
    expect(page.get_by_text(re.compile(r"目前名單：.*測試餐廳\.xlsx"))).to_be_visible()


def test_page_reload_returns_to_original_bundled_catalog(page):
    _import_catalog(page)
    page.reload()
    _expect_count(page, 10)
    expect(page.get_by_role("button", name="恢復內建名單", exact=True)).not_to_be_visible()
    expect(page.get_by_role("button", name="下載完整名單", exact=True)).to_be_visible()


def test_full_source_template_and_filtered_downloads_are_real_workbooks(page, tmp_path):
    original_path = tmp_path / "food.xlsx"
    _download_rows(page, "下載完整名單", original_path)
    assert original_path.read_bytes() == (ROOT / "data" / "food.xlsx").read_bytes()
    workbook = load_workbook(original_path, read_only=True)
    source_rows = [row for row in workbook.active.values if row[0]]
    workbook.close()
    assert source_rows[0][:11] == ("Restaurant", "Location", "Gener", "Price", "OnePerson", "Dating", "Gathering", "FastServe", "SlowEat", "Available time", "special available time")
    assert len(source_rows) == 11

    page.get_by_role("button", name="匯入 Excel 名單", exact=True).click()
    expect(page.get_by_role("dialog")).to_be_visible()
    template = _download_rows(page, "下載 Excel 範本", tmp_path / "template.xlsx")
    assert template[0] == EXCEL_HEADERS + ("Available time", "special available time")
    assert len(template) == 3
    assert all(row[0].startswith("示範・") for row in template[1:])
    page.keyboard.press("Escape")
    expect(page.get_by_role("dialog")).not_to_be_visible()

    _select_pills(page, "地點", ["小吃部"])
    _expect_count(page, 9)
    page.get_by_text("查看全部符合條件的餐廳（9 間）", exact=True).click()
    rows = _download_rows(page, "匯出篩選名單", tmp_path / "filtered.xlsx")
    assert rows[0] == EXCEL_HEADERS + ("Available time", "special available time")
    assert len(rows) == 10
    assert all(row[1] == "小吃部" for row in rows[1:])
    assert {row[0] for row in rows[1:]} == {"包子商行", "麥當勞", "比司多早午餐", "7-ELEVEN", "漢城異國美食", "蘇記牛肉麵", "一番快餐", "好富熟成黑咖哩.台式蓋飯", "珈名鮮茶"}
    assert sum(row[3] == 1.5 for row in rows[1:]) == 5
    assert "Available time" in rows[0] and "special available time" in rows[0]
    assert rows[1][rows[0].index("Available time")] == "7:30~19:30"
    assert rows[1][rows[0].index("special available time")] == "假日8:30~14:30"
    assert all(row[3] is None or 1 <= row[3] <= 5 for row in rows[1:])


def test_closed_restaurant_is_visible_only_on_request_and_never_recommended(page):
    content = _workbook([
        EXCEL_HEADERS,
        ["測試歇業店", "小吃部", "便當", 1, "O", "", "", "O", "", "", "關門了"],
        ["測試營業店", "清夜", "麵食", 2],
    ])
    dialog = _upload(page, content, "歇業測試名單.xlsx")
    expect(dialog).not_to_be_visible()
    _query(page, "測試歇業店")
    _expect_count(page, 0)
    page.get_by_text("更多選項", exact=True).click()
    checkbox = page.get_by_role("checkbox", name="包含已註記歇業店家", exact=True)
    page.get_by_text("包含已註記歇業店家", exact=True).click()
    expect(checkbox).to_be_checked()
    _expect_count(page, 1, available=0)
    expect(page.get_by_role("button", name="幫我決定這一餐", exact=True)).to_be_disabled()
    expect(page.get_by_text(re.compile("關門了")).first).to_be_visible()
    expect(page.locator(RESULT)).not_to_be_visible()
    page.get_by_role("button", name="全部取消", exact=True).click()
    _expect_count(page, 1)
    if not checkbox.is_visible():
        page.get_by_text("更多選項", exact=True).click()
    expect(checkbox).not_to_be_checked()


def test_multiple_pills_union_within_groups_and_toggle_off_without_dropdowns(page):
    content = _workbook([
        EXCEL_HEADERS,
        ["甲店", "Gongguan", "Bento", 1],
        ["乙店", "Shida", "Steak", 3],
        ["丙店", "Gongguan", "Steak", 2],
        ["丁店", "Liuzhangli", "Bento", 3],
        ["戊店", "Shida", "Cafe", 1],
        ["己店", "Gongguan", "Steak", 3],
    ])
    dialog = _upload(page, content, "多選名單.xlsx")
    expect(dialog).not_to_be_visible()
    _expect_count(page, 6)
    expect(page.get_by_role("combobox")).to_have_count(0)
    _select_pills(page, "地點", ["公館", "師大"])
    _expect_count(page, 5)
    _select_pills(page, "食物種類", ["便當", "牛排"])
    _expect_count(page, 4)
    _select_pills(page, "價格等級", ["🪙", "🪙🪙🪙"])
    _expect_count(page, 3)
    catalog = page.locator(".st-key-catalog_grid")
    for name in ("甲店", "乙店", "己店"):
        expect(catalog.get_by_role("heading", name=name, exact=True)).to_be_visible()
    for name in ("丙店", "丁店", "戊店"):
        expect(catalog.get_by_role("heading", name=name, exact=True)).to_have_count(0)

    _toggle_pill_off(page, "價格等級", "🪙")
    _expect_count(page, 2)
    _toggle_pill_off(page, "價格等級", "🪙🪙🪙")
    _expect_count(page, 4)
    page.get_by_role("button", name="全部取消", exact=True).click()
    _expect_count(page, 6)
    for label in ("地點", "食物種類", "價格等級", "用餐情境", "用餐節奏"):
        expect(page.get_by_role("toolbar", name=label, exact=True).get_by_role("button", pressed=True)).to_have_count(0)


def test_each_filter_group_has_select_all_and_cancel_buttons(page):
    groups = {
        "location_choices": ("地點", 2),
        "category_choices": ("食物種類", 18),
        "price_choices": ("價格等級", 5),
        "occasion_choice": ("用餐情境", 3),
        "pace_choice": ("用餐節奏", 2),
    }
    for key, (label, option_count) in groups.items():
        group = page.locator(f".st-key-group_{key}")
        group.get_by_role("button", name="全選", exact=True).click()
        selected = page.get_by_role("toolbar", name=label, exact=True).get_by_role("button", pressed=True)
        expect(selected).to_have_count(option_count)
        _wait_idle(page)
        group.get_by_role("button", name="取消", exact=True).click()
        expect(selected).to_have_count(0)
        _expect_count(page, 10)
    legend = page.get_by_role("list", name="金幣價格範圍說明", exact=True)
    for price_range in ("NT$0–100", "NT$100–200", "NT$200–300", "NT$300–500", "NT$500以上"):
        expect(legend).to_contain_text(price_range)
    price_buttons = page.get_by_role("toolbar", name="價格等級", exact=True)
    expect(price_buttons.get_by_role("button")).to_have_count(5)
    expect(price_buttons).not_to_contain_text("1.5")
    _select_pills(page, "價格等級", ["🪙"])
    _expect_count(page, 8)
    expect(page.locator(".st-key-catalog_grid").get_by_role("heading", name="7-ELEVEN", exact=True)).to_be_visible()
    _toggle_pill_off(page, "價格等級", "🪙")
    _select_pills(page, "價格等級", ["🪙🪙"])
    _expect_count(page, 7)
    expect(page.locator(".st-key-catalog_grid").get_by_role("heading", name="7-ELEVEN", exact=True)).to_be_visible()
    _select_pills(page, "價格等級", ["🪙"])
    _expect_count(page, 10)
    _toggle_pill_off(page, "價格等級", "🪙")
    _query(page, "7-ELEVEN")
    _expect_count(page, 1)
    page.get_by_role("button", name="幫我決定這一餐", exact=True).click()
    expect(page.locator(f"{RESULT} h2")).to_have_text("7-ELEVEN")
    expect(page.locator(RESULT).get_by_role("img", name="1.5 枚金幣", exact=True)).to_be_visible()


def test_all_catalog_pages_have_per_restaurant_comment_forms(page):
    previous = page.get_by_role("button", name="上一頁", exact=True)
    next_page = page.get_by_role("button", name="下一頁", exact=True)
    expect(previous).to_be_disabled()
    expect(page.get_by_role("heading", name="包子商行", exact=True)).to_be_visible()
    expect(page.get_by_test_id("stForm")).to_have_count(6)
    expect(page.get_by_role("textbox", name="評論內容", exact=True)).to_have_count(6)
    publish = page.get_by_role("button", name="發布評論", exact=True)
    expect(publish).to_have_count(6)
    for button in publish.all():
        expect(button).to_be_disabled()

    next_page.click()
    expect(page.get_by_role("heading", name="蘇記牛肉麵", exact=True)).to_be_visible()
    expect(page.get_by_role("heading", name="包子商行", exact=True)).to_have_count(0)
    expect(page.get_by_test_id("stForm")).to_have_count(4)
    expect(page.get_by_role("textbox", name="評論內容", exact=True)).to_have_count(4)
    expect(previous).to_be_enabled()
    previous.click()
    expect(page.get_by_role("heading", name="包子商行", exact=True)).to_be_visible()
    expect(previous).to_be_disabled()


def test_top_search_finds_names_outside_filters_and_closed_stores_before_the_catalog(page):
    _select_pills(page, "地點", ["清夜"])
    _expect_count(page, 1)
    _query(page, "麥當勞")
    results = page.locator(".st-key-search_results")
    expect(results.get_by_role("heading", name="麥當勞", exact=True)).to_be_visible()
    expect(results.get_by_role("textbox", name="您的暱稱", exact=True)).to_be_visible()
    expect(results.get_by_role("textbox", name="評論內容", exact=True)).to_be_visible()
    expect(results.get_by_role("button", name="發布評論", exact=True)).to_be_disabled()
    _expect_count(page, 0)
    results_bounds = results.bounding_box()
    filter_bounds = page.locator(".st-key-filters_panel").bounding_box()
    assert results_bounds is not None and filter_bounds is not None
    assert results_bounds["y"] < filter_bounds["y"]

    _query(page, "")
    content = _workbook([
        EXCEL_HEADERS,
        ["測試歇業店", "小吃部", "便當", 1, "O", "", "", "O", "", "", "關門了"],
        ["測試營業店", "清夜", "麵食", 2],
    ])
    dialog = _upload(page, content, "歇業測試名單.xlsx")
    expect(dialog).not_to_be_visible()
    _query(page, "測試歇業店")
    expect(results.get_by_role("heading", name="測試歇業店", exact=True)).to_be_visible()
    expect(results).to_contain_text("關門了")
    expect(results.get_by_role("textbox", name="評論內容", exact=True)).to_be_visible()
    expect(page.get_by_role("button", name="幫我決定這一餐", exact=True)).to_be_disabled()
    _query(page, "")
    expect(results).to_have_count(0)
    _expect_count(page, 1)
    page.get_by_role("button", name="恢復內建名單", exact=True).click()
    _expect_count(page, 10)


def test_one_business_can_have_independent_search_recommendation_and_catalog_half_star_forms(page):
    _query(page, "麥當勞")
    _expect_count(page, 1)
    results = page.locator(".st-key-search_results")
    expect(results.get_by_role("heading", name="麥當勞", exact=True)).to_be_visible()
    page.get_by_role("button", name="幫我決定這一餐", exact=True).click()
    recommendation = page.locator(".st-key-pick_panel")
    expect(page.locator(f"{RESULT} h2")).to_have_text("麥當勞")
    expect(page.get_by_test_id("stForm")).to_have_count(3)
    for context in (results, recommendation, page.locator(".st-key-catalog_grid")):
        expect(context.get_by_role("textbox", name="您的暱稱", exact=True)).to_be_visible()
        expect(context.get_by_role("textbox", name="評論內容", exact=True)).to_be_visible()
        expect(context.get_by_role("button", name="發布評論", exact=True)).to_be_disabled()
    rating = recommendation.get_by_role("slider")
    expect(rating).to_have_value("3")
    expect(rating).to_have_attribute("step", "0.5")
    rating.press("Home")
    rating.press("ArrowRight")
    expect(rating).to_have_value("1.5")
    expect(results.get_by_role("slider")).to_have_value("3")
    recommendation.get_by_role("textbox", name="評論內容", exact=True).fill("測試草稿，不會發布")
    expect(results.get_by_role("textbox", name="評論內容", exact=True)).to_have_value("")


def test_top_search_pages_expose_comment_forms_for_every_matching_store(page):
    content = _workbook([EXCEL_HEADERS, *[
        [f"搜尋測試{index}店", "小吃部", "便當", 1] for index in range(6)
    ]])
    dialog = _upload(page, content, "搜尋分頁名單.xlsx")
    expect(dialog).not_to_be_visible()
    _query(page, "搜尋測試")
    results = page.locator(".st-key-search_results")
    expect(results.get_by_test_id("stForm")).to_have_count(3)
    expect(results.get_by_role("heading", name="搜尋測試0店", exact=True)).to_be_visible()
    previous = results.get_by_role("button", name="上一頁搜尋結果", exact=True)
    next_page = results.get_by_role("button", name="下一頁搜尋結果", exact=True)
    expect(previous).to_be_disabled()
    next_page.click()
    expect(results.get_by_role("heading", name="搜尋測試0店", exact=True)).to_have_count(0)
    expect(results.get_by_test_id("stForm")).to_have_count(3)
    expect(previous).to_be_enabled()
    previous.click()
    expect(results.get_by_role("heading", name="搜尋測試0店", exact=True)).to_be_visible()
    _query(page, "搜尋測試0店")
    expect(results.get_by_test_id("stForm")).to_have_count(1)


def test_search_clear_icon_is_inside_the_field_and_clears_unsent_text_preserving_filters(page):
    query = page.get_by_role("textbox", name="搜尋店家", exact=True)
    clear = page.get_by_role("button", name="清除搜尋", exact=True)
    expect(clear).to_be_visible()
    field_bounds = query.bounding_box()
    button_bounds = clear.bounding_box()
    assert field_bounds is not None and button_bounds is not None
    assert field_bounds["x"] + field_bounds["width"] * .7 <= button_bounds["x"]
    assert button_bounds["x"] + button_bounds["width"] <= field_bounds["x"] + field_bounds["width"] + 1
    assert field_bounds["y"] - 1 <= button_bounds["y"]
    assert button_bounds["y"] + button_bounds["height"] <= field_bounds["y"] + field_bounds["height"] + 1
    if page.viewport_size["width"] <= 640:
        assert button_bounds["width"] >= 40 and button_bounds["height"] >= 40

    _select_pills(page, "地點", ["小吃部"])
    _select_pills(page, "食物種類", ["速食"])
    _query(page, "麥當勞")
    expect(page.locator(".st-key-search_results").get_by_role("heading", name="麥當勞", exact=True)).to_be_visible()
    query.fill("尚未按 Enter 的文字")
    clear.click()
    expect(query).to_have_value("")
    expect(page.locator(".st-key-search_results")).to_have_count(0)
    _expect_count(page, 1)
    for label, name in (("地點", "小吃部"), ("食物種類", "速食")):
        expect(page.get_by_role("toolbar", name=label, exact=True).get_by_role("button", name=name, exact=True)).to_have_attribute("aria-pressed", "true")

    _query(page, "麥當勞")
    query.focus()
    query.press("Tab")
    expect(clear).to_be_focused()
    clear.press("Enter")
    expect(query).to_have_value("")
    _expect_count(page, 1)


def test_homophone_typo_directly_shows_a_real_store_and_review_form_in_search_results(page):
    _select_pills(page, "地點", ["清夜"])
    query = _query(page, "麥當牢")
    expect(query).to_have_value("麥當牢")
    results = page.locator(".st-key-search_results")
    expect(results).to_have_count(1)
    expect(results.get_by_text("找到 1 間讀音或字詞相近的店家", exact=True)).to_be_visible()
    expect(results.get_by_role("heading", name="麥當勞", exact=True)).to_be_visible()
    expect(results.get_by_test_id("stForm")).to_have_count(1)
    expect(results.get_by_role("textbox", name="您的暱稱", exact=True)).to_be_visible()
    expect(results.get_by_role("textbox", name="評論內容", exact=True)).to_be_visible()
    expect(results.get_by_role("slider")).to_have_attribute("step", "0.5")
    expect(results.get_by_role("button", name="發布評論", exact=True)).to_be_disabled()
    expect(results.get_by_text(re.compile("找到 0 間|找不到完全符合"))).to_have_count(0)
    expect(page.get_by_role("heading", name="你是不是想找這些店家？", exact=True)).to_have_count(0)
    expect(page.locator(".st-key-search_suggestions")).to_have_count(0)
    _expect_count(page, 0)

    _query(page, "ZXQ987654321外太空不存在的店")
    expect(results.get_by_test_id("stForm")).to_have_count(0)
    expect(results.get_by_text("找不到相近的店家，請試試其他店名或較短的關鍵字。", exact=True)).to_be_visible()
    page.get_by_role("button", name="清除搜尋", exact=True).click()
    expect(query).to_have_value("")
    expect(results).to_have_count(0)
    _expect_count(page, 1)


def test_typing_short_homophone_and_exact_names_updates_search_without_enter_or_blur(page):
    query = page.get_by_role("textbox", name="搜尋店家", exact=True)
    query.fill("麥當牢")
    results = page.locator(".st-key-search_results")
    expect(results.get_by_role("heading", name="麥當勞", exact=True)).to_be_visible()
    expect(query).to_be_focused()
    expect(query).to_have_value("麥當牢")
    expect(results.get_by_role("textbox", name="評論內容", exact=True)).to_be_visible()
    expect(results.get_by_role("slider")).to_have_attribute("step", "0.5")

    query.fill("")
    expect(results).to_have_count(0)
    expect(query).to_be_focused()
    query.fill("麥當勞")
    expect(results.get_by_role("heading", name="麥當勞", exact=True)).to_be_visible()
    expect(query).to_be_focused()
    expect(results.get_by_role("button", name="發布評論", exact=True)).to_be_disabled()


def test_rapid_typing_and_clear_cancel_pending_live_search_updates(page):
    query = page.get_by_role("textbox", name="搜尋店家", exact=True)
    query.fill("包子傷行")
    query.fill("麥當牢")
    expect(page.locator(".st-key-search_results").get_by_role(
        "heading", name="麥當勞", exact=True
    )).to_be_visible()
    expect(query).to_be_focused()
    expect(query).to_have_value("麥當牢")

    query.fill("麥當勞")
    page.get_by_role("button", name="清除搜尋", exact=True).click()
    expect(query).to_have_value("")
    expect(page.locator(".st-key-search_results")).to_have_count(0)
    # A cancelled 350 ms callback must stay cancelled after its original deadline.
    page.wait_for_timeout(600)
    expect(query).to_have_value("")
    expect(page.locator(".st-key-search_results")).to_have_count(0)

    query.fill("包子傷行")
    query.fill("麥當勞")
    results = page.locator(".st-key-search_results")
    expect(results.get_by_role("heading", name="麥當勞", exact=True)).to_be_visible()
    page.wait_for_timeout(600)
    expect(query).to_be_focused()
    expect(query).to_have_value("麥當勞")
    expect(results.get_by_role("heading", name="包子商行", exact=True)).to_have_count(0)

"""Verify real workbook boundaries, exported text, and import behavior."""

from dataclasses import replace
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook, load_workbook

from food_picker.catalog import demo_restaurants
from food_picker.defaults import load_default_catalog
from food_picker.excel import (
    EXCEL_HEADERS,
    LEGACY_HEADERS,
    SCHEDULE_HEADERS,
    MAX_UNCOMPRESSED_BYTES,
    MAX_UPLOAD_BYTES,
    MAX_ZIP_ENTRIES,
    build_template,
    export_restaurants,
    parse_excel,
)
from food_picker.models import ValidationError


def workbook_bytes(rows, second_sheet=None):
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    if second_sheet is not None:
        other = workbook.create_sheet("其他資料")
        for row in second_sheet:
            other.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def test_template_uses_full_schema_and_clearly_fictional_examples():
    content = build_template()
    result = parse_excel(content, "餐廳.XLSX")
    assert len(result.restaurants) == 2
    assert not result.warnings
    assert all(item.name.startswith("示範・") and "虛構" in item.note for item in result.restaurants)
    assert [item.location for item in result.restaurants] == ["小吃部", "清夜"]
    assert [item.price_level for item in result.restaurants] == [1.5, 3]
    assert result.restaurants[0].opening_hours == "7:00~19:00"
    assert result.restaurants[0].special_opening_hours == "假日營業時間另行公告"
    assert all(item.price is None for item in result.restaurants)
    assert result.restaurants[0].one_person and result.restaurants[0].fast_serve
    assert result.restaurants[1].dating and result.restaurants[1].gathering and result.restaurants[1].slow_eat
    with BytesIO(content) as stream:
        workbook = load_workbook(stream)
        sheet = workbook.active
        assert tuple(cell.value for cell in sheet[1]) == EXCEL_HEADERS + SCHEDULE_HEADERS
        assert sheet.freeze_panes == "A2"
        assert sheet["A1"].font.bold
        assert sheet["E1"].comment
        workbook.close()


def test_only_first_sheet_is_imported():
    content = workbook_bytes(
        [LEGACY_HEADERS, ["第一間", "校內", "中式", 80]],
        [LEGACY_HEADERS, ["不應出現", "校內", "中式", 90]],
    )
    assert [item.name for item in parse_excel(content, "餐廳.xlsx").restaurants] == ["第一間"]


def test_empty_first_sheet_does_not_fall_back_to_second():
    content = workbook_bytes([], [LEGACY_HEADERS, ["第二間", "校內"]])
    with pytest.raises(ValidationError, match="沒有資料"):
        parse_excel(content, "餐廳.xlsx")


def test_import_keeps_excel_row_numbers_and_domain_warnings():
    content = workbook_bytes(
        [
            [],
            LEGACY_HEADERS,
            ["可用", "校內", "中式", 80],
            ["過遠", "校外", "中式", 100, 2.1],
            ["未確認距離", "校外", "中式", 100],
        ]
    )
    result = parse_excel(content, "餐廳.xlsx")
    assert [item.name for item in result.restaurants] == ["可用", "未確認距離"]
    assert any("4" in warning and "2 公里" in warning for warning in result.warnings)
    assert any("5" in warning and "距離" in warning for warning in result.warnings)


def test_formulas_are_not_replaced_by_cached_values():
    content = workbook_bytes(
        [
            LEGACY_HEADERS,
            ["可用", "校內", "中式", 80],
            ["公式價錢", "校內", "中式", "=SUM(40,40)"],
            ["=CONCAT(\"公式名稱\")", "校內", "中式", 90],
        ]
    )
    result = parse_excel(content, "餐廳.xlsx")
    assert [item.name for item in result.restaurants] == ["可用"]
    assert result.warnings


@pytest.mark.parametrize("filename", ["餐廳.xls", "餐廳.xlsm", "餐廳.csv", "餐廳.xlsx.zip"])
def test_unsupported_extensions_explain_xlsx_conversion(filename):
    with pytest.raises(ValidationError, match=".xlsx"):
        parse_excel(b"content", filename)


@pytest.mark.parametrize("content", [b"not a workbook", b"PK\x03\x04", b"<html>error</html>"])
def test_corrupt_archives_have_actionable_error(content):
    with pytest.raises(ValidationError, match="未加密"):
        parse_excel(content, "餐廳.xlsx")


def test_empty_and_oversized_uploads():
    with pytest.raises(ValidationError, match="空"):
        parse_excel(b"", "餐廳.xlsx")
    with pytest.raises(ValidationError, match="5 MB"):
        parse_excel(b"0" * (MAX_UPLOAD_BYTES + 1), "餐廳.xlsx")


def test_non_excel_zip_is_not_treated_as_workbook():
    buffer = BytesIO()
    with ZipFile(buffer, "w") as archive:
        archive.writestr("hello.txt", "Hello")
    with pytest.raises(ValidationError, match="有效"):
        parse_excel(buffer.getvalue(), "餐廳.xlsx")


def test_incomplete_excel_archive_is_reported_cleanly():
    buffer = BytesIO()
    with ZipFile(buffer, "w") as archive:
        archive.writestr("[Content_Types].xml", "<broken")
    with pytest.raises(ValidationError, match="有效"):
        parse_excel(buffer.getvalue(), "餐廳.xlsx")


@pytest.mark.parametrize("unsafe_path", ["../outside.xml", "/outside.xml", "xl/../../outside.xml", "..\\outside.xml"])
def test_archive_paths_cannot_escape_workbook(unsafe_path):
    buffer = BytesIO()
    with ZipFile(buffer, "w") as archive:
        archive.writestr("[Content_Types].xml", "placeholder")
        archive.writestr(unsafe_path, "placeholder")
    with pytest.raises(ValidationError, match="有效"):
        parse_excel(buffer.getvalue(), "餐廳.xlsx")


def test_archive_entry_count_is_bounded():
    buffer = BytesIO()
    with ZipFile(buffer, "w") as archive:
        for number in range(MAX_ZIP_ENTRIES + 1):
            archive.writestr(f"part-{number}.xml", "")
    with pytest.raises(ValidationError, match="過多"):
        parse_excel(buffer.getvalue(), "餐廳.xlsx")


def test_decompression_size_is_bounded_before_loading():
    buffer = BytesIO()
    with ZipFile(buffer, "w", ZIP_DEFLATED) as archive:
        archive.writestr("large.xml", b"0" * (MAX_UNCOMPRESSED_BYTES + 1))
    assert len(buffer.getvalue()) < MAX_UPLOAD_BYTES
    with pytest.raises(ValidationError, match="32 MB"):
        parse_excel(buffer.getvalue(), "餐廳.xlsx")


@pytest.mark.parametrize("row,column,expected", [(10_001, 1, "10,000"), (1, 33, "32 欄")])
def test_worksheet_dimensions_are_bounded(row, column, expected):
    workbook = Workbook()
    workbook.active.append(LEGACY_HEADERS)
    workbook.active.append(["可用", "校內"])
    workbook.active.cell(row, column, "超出範圍")
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    with pytest.raises(ValidationError, match=expected):
        parse_excel(buffer.getvalue(), "餐廳.xlsx")


def test_domain_restaurant_limit_still_applies_to_real_workbook():
    content = workbook_bytes([LEGACY_HEADERS] + [[f"店家{number}", "校內"] for number in range(5001)])
    with pytest.raises(ValidationError, match="5,000"):
        parse_excel(content, "餐廳.xlsx")


def test_export_contains_only_requested_restaurants():
    selected = demo_restaurants()[2:4]
    result = parse_excel(export_restaurants(selected), "已篩選.xlsx")
    assert [item.name for item in result.restaurants] == [item.name for item in selected]


def test_export_roundtrip_preserves_all_bundled_catalog_fields():
    original = load_default_catalog().restaurants
    imported = parse_excel(export_restaurants(original), "完整餐廳.xlsx").restaurants
    assert len(original) == len(imported) == 10
    assert imported == original


def test_export_preserves_legacy_amount_distance_address_and_area():
    original = demo_restaurants()
    imported = parse_excel(export_restaurants(original), "舊格式匯出.xlsx").restaurants
    for before, after in zip(original, imported, strict=True):
        assert (after.name, after.area, after.category, after.price, after.distance, after.address, after.note) == (
            before.name, before.area, before.category, before.price, before.distance, before.address, before.note
        )
        assert after.price_level is None


def test_export_preserves_rich_and_legacy_records_in_the_same_workbook():
    original = [load_default_catalog().restaurants[0], demo_restaurants()[4]]
    imported = parse_excel(export_restaurants(original), "混合名單.xlsx").restaurants
    assert len(imported) == 2
    assert imported[0] == original[0]
    assert imported[1].price == original[1].price
    assert imported[1].distance == original[1].distance
    assert imported[1].area == original[1].area


def test_exported_untrusted_strings_are_literal_cells():
    restaurant = replace(
        demo_restaurants()[0],
        name="=HYPERLINK(\"https://example.com\")",
        category="+SUM(1,1)",
        address="-1+2",
        note="@SUM(1,1)",
    )
    workbook = load_workbook(BytesIO(export_restaurants([restaurant])), data_only=False)
    sheet = workbook.active
    for coordinate, expected in (
        ("A2", restaurant.name),
        ("C2", restaurant.category),
        ("O2", restaurant.address),
        ("K2", restaurant.note),
    ):
        assert sheet[coordinate].value == expected
        assert sheet[coordinate].data_type == "s"
    workbook.close()


def test_empty_export_retains_headers_for_further_editing():
    workbook = load_workbook(BytesIO(export_restaurants([])))
    assert workbook.active.max_row == 1
    assert tuple(cell.value for cell in workbook.active[1]) == EXCEL_HEADERS
    workbook.close()

"""Read and write bounded Excel workbooks without running spreadsheet formulas."""

from __future__ import annotations

from collections.abc import Sequence
from io import BytesIO
from pathlib import PurePosixPath
from zipfile import BadZipFile, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from food_picker.catalog import parse_rows
from food_picker.models import ParseResult, Restaurant, ValidationError

EXCEL_HEADERS = (
    "Restaurant",
    "Location",
    "Genre",
    "Price",
    "OnePerson",
    "Dating",
    "Gathering",
    "FastServe",
    "SlowEat",
    "Coordinates",
    "Comments",
)
LEGACY_HEADERS = (
    "餐廳名稱",
    "區域",
    "類型",
    "均消",
    "距離（公里）",
    "地址",
    "備註",
)
MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 32 * 1024 * 1024
MAX_ZIP_ENTRIES = 256
MAX_WORKSHEET_ROWS = 10_000
MAX_WORKSHEET_COLUMNS = 32

_UNREADABLE = "無法讀取 Excel，請確認檔案未加密，並以 Excel 另存為有效的 .xlsx 格式。"


def _check_archive(content: bytes) -> None:
    """Check container limits before openpyxl expands any workbook parts."""
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ZIP_ENTRIES:
                raise ValidationError("Excel 內含過多檔案，請移除不需要的圖片或工作表後重試。")
            if sum(entry.file_size for entry in entries) > MAX_UNCOMPRESSED_BYTES:
                raise ValidationError("Excel 解壓縮後超過 32 MB，請移除圖片或分割檔案後重試。")
            for entry in entries:
                path = PurePosixPath(entry.filename.replace("\\", "/"))
                if path.is_absolute() or ".." in path.parts or entry.flag_bits & 1:
                    raise ValidationError(_UNREADABLE)
            if "[Content_Types].xml" not in archive.namelist():
                raise ValidationError(_UNREADABLE)
    except (BadZipFile, OSError, ValueError) as error:
        if isinstance(error, ValidationError):
            raise
        raise ValidationError(_UNREADABLE) from error


def parse_excel(content: bytes, filename: str) -> ParseResult:
    """Read only the first worksheet, keeping formulas visible to row validation."""
    if not filename.lower().endswith(".xlsx"):
        raise ValidationError("請選擇 .xlsx 格式；舊版 .xls 請先以 Excel 另存為 .xlsx。")
    if not content:
        raise ValidationError("檔案是空的，請重新選擇。")
    if len(content) > MAX_UPLOAD_BYTES:
        raise ValidationError("檔案超過 5 MB，請刪除不需要的工作表或圖片後重試。")
    _check_archive(content)

    workbook = None
    try:
        workbook = load_workbook(
            BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
        if not workbook.worksheets:
            raise ValidationError("Excel 沒有工作表，請加入餐廳資料後重試。")
        sheet = workbook.worksheets[0]
        if sheet.max_row and sheet.max_row > MAX_WORKSHEET_ROWS:
            raise ValidationError("第一個工作表超過 10,000 列，請刪除多餘的空白列或分割檔案。")
        if sheet.max_column and sheet.max_column > MAX_WORKSHEET_COLUMNS:
            raise ValidationError("第一個工作表超過 32 欄，請保留餐廳資料欄位後重試。")
        # Ignore incorrect worksheet dimensions. A fixed column bound and explicit
        # row count keep malformed files from allocating an enormous blank grid.
        sheet.reset_dimensions()
        rows = []
        for number, row in enumerate(
            sheet.iter_rows(max_col=MAX_WORKSHEET_COLUMNS + 1, values_only=True), 1
        ):
            if number > MAX_WORKSHEET_ROWS:
                raise ValidationError("第一個工作表超過 10,000 列，請刪除多餘的空白列或分割檔案。")
            if row[MAX_WORKSHEET_COLUMNS] is not None:
                raise ValidationError("第一個工作表超過 32 欄，請保留餐廳資料欄位後重試。")
            rows.append(list(row[:MAX_WORKSHEET_COLUMNS]))
    except ValidationError:
        raise
    except Exception as error:
        # openpyxl exposes several XML/ZIP exceptions for different corrupt parts.
        raise ValidationError(_UNREADABLE) from error
    finally:
        if workbook is not None:
            workbook.close()
    return parse_rows(rows)


def _write_workbook(headers: Sequence[str], rows: Sequence[Sequence[object]]) -> bytes:
    """Write literal text cells, never spreadsheet formulas or external links."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "餐廳名單"
    sheet.append(tuple(headers))
    for row in rows:
        sheet.append(tuple(row))
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"

    for cell in sheet[1]:
        cell.font = Font(name="Microsoft JhengHei", bold=True, color="253224")
        cell.fill = PatternFill("solid", fgColor="E5EBDD")
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 25
    widths = (30, 21, 22, 12, 14, 12, 14, 14, 14, 43, 54, 16, 14, 18, 38)
    for index, width in enumerate(widths[:len(headers)], 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet["B1"].comment = Comment("保留原始地區代碼，例如 Gongguan、School；不會自動判定清大周邊距離。", "食刻決定")
    sheet["D1"].comment = Comment("價格級距：請填 1 至 5；這不是新臺幣金額，不知道時可留白。", "食刻決定")
    for column in "EFGHI":
        sheet[f"{column}1"].comment = Comment("適合此情境時填 O，不適合或未知時留白。", "食刻決定")
    sheet["J1"].comment = Comment("緯度, 經度，例如 25.0182, 121.5310；不知道時可留白。", "食刻決定")
    sheet["K1"].comment = Comment("店家備註會原樣保留；包含「關門了」等歇業提醒的店家不會被推薦。", "食刻決定")
    output = BytesIO()
    try:
        workbook.save(output)
    finally:
        workbook.close()
    return output.getvalue()


def export_restaurants(restaurants: Sequence[Restaurant]) -> bytes:
    """Export all catalog fields, preserving raw region/genre codes and old data."""
    legacy_extensions = any(
        restaurant.area in ("campus", "nearby")
        or restaurant.price is not None
        or restaurant.distance is not None
        or restaurant.address
        for restaurant in restaurants
    )
    headers = EXCEL_HEADERS + (("區域", "均消", "距離（公里）", "地址") if legacy_extensions else ())
    rows = []
    for restaurant in restaurants:
        coordinates = (
            f"{restaurant.latitude!r}, {restaurant.longitude!r}"
            if restaurant.latitude is not None and restaurant.longitude is not None
            else ""
        )
        values = (
            restaurant.name,
            restaurant.location,
            restaurant.genre or restaurant.category,
            restaurant.price_level,
            *("O" if value else "" for value in (
                restaurant.one_person, restaurant.dating, restaurant.gathering,
                restaurant.fast_serve, restaurant.slow_eat,
            )),
            coordinates,
            restaurant.note,
        )
        if legacy_extensions:
            area = {"campus": "清大校內", "nearby": "校園周邊"}.get(restaurant.area, "")
            values += (area, restaurant.price, restaurant.distance, restaurant.address)
        rows.append(values)
    return _write_workbook(headers, rows)


def build_template() -> bytes:
    """Return the eleven-column schema with two explicitly fictional examples."""
    return _write_workbook(EXCEL_HEADERS, [
        ("示範・一人食堂", "Gongguan", "Bento", 1, "O", "", "", "O", "", "", "虛構示範資料，非真實店家"),
        ("示範・聚餐小館", "School", "Chinese", 3, "", "O", "O", "", "O", "", "虛構示範資料，非真實店家"),
    ])
